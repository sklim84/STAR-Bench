"""AML 벤치마크 데이터셋 Excel 생성 스크립트.

아래 4개 시트를 하나의 Excel 파일로 생성한다.

    Sheet 1 – 도구 정의:
        20개 에이전트 도구의 ①정답/②유사/③미끼 역할 분류 및 파라미터 스펙.
    Sheet 2 – 평가 시나리오:
        22개 대표 시나리오 (상황 설정·사용자 질문·정답·근거·모델 사고과정).
    Sheet 3 – 벤치마크 케이스:
        dataset/*.json 에서 읽어온 1,258개 실행 케이스 전체 (24개 카테고리).
    Sheet 4 – 케이스 요약:
        카테고리별 케이스 수 집계.

실행:
    python _paper/benchmarks/create_benchmark_dataset.py
"""

from __future__ import annotations

import json
import os
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

# ---------------------------------------------------------------------------
# 공통 스타일
# ---------------------------------------------------------------------------

_THIN = Side(style="thin", color="BDBDBD")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_WRAP_TOP = Alignment(wrap_text=True, vertical="top")
_CENTER_MID = Alignment(horizontal="center", vertical="center", wrap_text=True)

_HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
_HEADER_FILL = PatternFill(start_color="2E75B6", end_color="2E75B6", fill_type="solid")
_SUB_HEADER_FILL = PatternFill(start_color="8EA9DB", end_color="8EA9DB", fill_type="solid")
_BOLD = Font(bold=True)

# ---------------------------------------------------------------------------
# ① 도구 정의 데이터
# ---------------------------------------------------------------------------
# 각 항목: (sub_domain, task, func_name, func_desc, [(required_flag, prop, type, desc), ...])

_TOOLS: list[tuple] = [

    # ===================================================================
    # Sub domain 1: 이상거래 통계 및 현황 조회
    # ===================================================================

    # --- get_statistics (① 정답 도구) ---
    (
        "이상거래 통계 및 현황 조회",
        "전체 이상거래 요약 통계 조회 (① 정답 도구)",
        "get_statistics",
        "HOFINET 전체 거래 데이터의 요약 통계를 조회한다. "
        "전체 거래 건수, 이상거래 건수·비율, 이상거래 유형별 분포(1~7)를 반환한다.",
        [
            ("optional", "(없음)", "-", "파라미터 없이 호출. 전체 기간·전체 유형 집계 반환."),
        ],
    ),

    # --- get_fraud_type_summary (① 정답 도구) ---
    (
        "이상거래 통계 및 현황 조회",
        "이상거래 유형별 상세 현황 조회 (① 정답 도구)",
        "get_fraud_type_summary",
        "이상거래 유형(코드 1~7)을 지정하면 해당 유형의 건수·금액 통계, "
        "상위 금융회사, 샘플 날짜를 반환한다. "
        "코드 매핑: 1=자금세탁, 2=신규거래처, 3=대포통장, 4=보이스피싱, "
        "5=불법도박, 6=유사수신, 7=기타",
        [
            ("required", "fraud_type", "integer", "이상거래유형 코드 (1~7). 예: 4=보이스피싱"),
            ("optional", "bank_id", "integer",
             "출금금융회사일련번호 필터 (선택). 지정 시 해당 금융회사 거래만 집계."),
        ],
    ),

    # --- get_account_profile (② 유사 기능 도구) ---
    (
        "이상거래 통계 및 현황 조회",
        "특정 계좌의 거래 통계 프로파일 조회 (② 유사 기능 도구)",
        "get_account_profile",
        "특정 계좌의 거래 통계 프로파일을 조회한다. "
        "총 거래 건수·금액, 이상거래 건수·비율, 주요 거래 시간대, "
        "주 사용 매체, 상위 거래 상대 계좌 5개를 반환한다. "
        "전체 통계가 아닌 단일 계좌 단위 집계이므로 유형별 전체 현황 파악에는 부적합하다.",
        [
            ("required", "account_id", "integer",
             "조회할 계좌 번호 (출금계좌일련번호 기준). 예: 1234567890"),
        ],
    ),

    # --- compare_periods (① 정답 도구) ---
    (
        "이상거래 통계 및 현황 조회",
        "두 기간 이상거래 통계 비교 (① 정답 도구)",
        "compare_periods",
        "두 기간(YYYYMMDD 범위)의 거래·이상거래 통계를 비교하고 주요 변화를 반환한다. "
        "각 기간의 거래 건수·이상거래 건수·평균 금액·이상거래 비율과 "
        "기간 간 delta(증감율·증감폭)를 함께 반환한다.",
        [
            ("required", "period1_start", "integer",
             "비교 기간 1 시작일 (YYYYMMDD). 예: 20231001"),
            ("required", "period1_end", "integer",
             "비교 기간 1 종료일 (YYYYMMDD). 예: 20231231"),
            ("required", "period2_start", "integer",
             "비교 기간 2 시작일 (YYYYMMDD). 예: 20240101"),
            ("required", "period2_end", "integer",
             "비교 기간 2 종료일 (YYYYMMDD). 예: 20240331"),
        ],
    ),

    # --- get_institution_report (① 정답 도구) ---
    (
        "이상거래 통계 및 현황 조회",
        "금융회사 종합 현황 보고 (① 정답 도구)",
        "get_institution_report",
        "특정 금융회사의 종합 현황을 보고한다. "
        "출금 기준 총 거래 건수·금액, 이상거래 건수·비율, "
        "상위 거래 상대 기관, 이상거래 유형 분포를 반환한다.",
        [
            ("required", "bank_id", "integer",
             "조회할 금융회사 일련번호 (출금금융회사일련번호 기준). 예: 134"),
        ],
    ),

    # --- query_transactions (③ 인접·미끼 도구) ---
    (
        "이상거래 통계 및 현황 조회",
        "SQL 기반 거래 데이터 원본 조회 (③ 인접·미끼 도구)",
        "query_transactions",
        "HOFINET 데이터베이스에 SQL 쿼리를 실행하여 거래 원본 레코드를 조회한다. "
        "테이블명 hofinet, 컬럼은 거래일자·거래시간대·출금/입금금융회사일련번호·"
        "출금/입금계좌일련번호·자금구분·매체구분·거래금액·이상거래여부·이상거래유형·이상거래설명이다. "
        "집계 통계가 아닌 row 단위 반환이므로 통계 현황 확인 목적에는 부적합하다.",
        [
            ("required", "sql", "string",
             "실행할 SELECT SQL. hofinet 테이블 대상. "
             "예: SELECT * FROM hofinet WHERE 이상거래여부=1 LIMIT 10"),
        ],
    ),

    # ===================================================================
    # Sub domain 2: AML 탐지·분석 및 보고서 작성
    # ===================================================================

    # --- analyze_network (① 정답 도구) ---
    (
        "AML 탐지·분석 및 보고서 작성",
        "계좌 거래 네트워크 구조 분석 (① 정답 도구)",
        "analyze_network",
        "특정 계좌를 중심으로 거래 네트워크를 분석한다. "
        "연결된 계좌 수, 거래 횟수, 이상거래 관련 여부, N-hop 심층 탐색을 지원한다. "
        "계좌 간 연결 구조와 규모를 파악하는 데 적합하다.",
        [
            ("required", "account_id", "integer",
             "분석할 계좌 번호 (출금계좌일련번호). 예: 1234567890"),
            ("optional", "hops", "integer",
             "탐색 범위 (1~5). 기본값 1. 3 이상은 Memgraph 필요."),
        ],
    ),

    # --- detect_aml_patterns (② 유사 기능 도구) ---
    (
        "AML 탐지·분석 및 보고서 작성",
        "AML 특화 패턴 탐지 (② 유사 기능 도구)",
        "detect_aml_patterns",
        "Memgraph 그래프 DB를 활용하여 AML 특화 패턴을 탐지한다. "
        "순환거래(ring), 다단계 레이어링(layering), 대포통장(funnel) 패턴을 탐지하거나 "
        "두 계좌 간 최단경로(shortest_path) 조회, 계좌 위험도 점수(risk_score) 산출을 수행한다. "
        "네트워크 구조가 아닌 특정 AML 패턴 유형 탐지가 목적이다.",
        [
            ("required", "pattern_type", "string",
             "탐지 패턴 유형: ring / layering / funnel / shortest_path / risk_score"),
            ("optional", "account_id", "integer",
             "분석 대상 계좌 번호 (risk_score 시 필수)"),
            ("optional", "account_a", "integer",
             "출발 계좌 번호 (shortest_path 시 필수)"),
            ("optional", "account_b", "integer",
             "도착 계좌 번호 (shortest_path 시 필수)"),
        ],
    ),

    # --- rank_risky_transactions (① 정답 도구) ---
    (
        "AML 탐지·분석 및 보고서 작성",
        "위험도 상위 거래 일괄 랭킹 (① 정답 도구)",
        "rank_risky_transactions",
        "저장된 XGBoost 모델로 샘플 N건을 일괄 예측하고 이상거래 확률 기준 상위 K건을 반환한다. "
        "단건 예측(predict_fraud)과 달리 대량 거래에서 위험도 높은 거래를 자동으로 발굴하는 데 사용한다.",
        [
            ("optional", "sample_size", "integer",
             "일괄 예측할 샘플 건수 (기본값 1000, 최대 5000). 예: 500"),
            ("optional", "top_k", "integer",
             "반환할 상위 K건 (기본값 20, 최대 100). 예: 10"),
        ],
    ),

    # --- predict_fraud (③ 인접·미끼 도구) ---
    (
        "AML 탐지·분석 및 보고서 작성",
        "XGBoost 모델 기반 이상거래 확률 예측 (③ 인접·미끼 도구)",
        "predict_fraud",
        "학습된 XGBoost 모델로 단건 거래의 이상거래 확률을 예측한다. "
        "거래 6개 속성(시간대·출금기관·입금기관·자금구분·매체구분·금액)을 입력하면 0~1 확률을 반환한다. "
        "계좌·네트워크 분석이 아닌 개별 거래 단위 ML 예측이므로 "
        "네트워크·패턴 분석 목적에는 부적합하다.",
        [
            ("required", "거래시간대", "integer",
             "3시간 단위 시간대 (0, 3, 6, 9, 12, 15, 18, 21 중 하나)"),
            ("required", "출금금융회사일련번호", "integer", "출금 금융회사 일련번호"),
            ("required", "입금금융회사일련번호", "integer", "입금 금융회사 일련번호"),
            ("required", "자금구분", "integer", "자금구분 코드 (0, 1, 3, 4 중 하나)"),
            ("required", "매체구분", "integer", "매체구분 코드 (1~7 중 하나)"),
            ("required", "거래금액", "integer", "원 단위 거래 금액 (양의 정수)"),
        ],
    ),

    # --- generate_str (① 정답 도구) ---
    (
        "AML 탐지·분석 및 보고서 작성",
        "의심거래보고서(STR) 공식 양식 작성 (① 정답 도구)",
        "generate_str",
        "분석 결과를 기반으로 의심거래보고서(STR) 공식 양식(I~VII 섹션)에 맞게 작성한다. "
        "STR 작성 전에 반드시 query_transactions로 관련 거래 데이터를 먼저 조회해야 한다. "
        "transactions(거래 레코드), fraud_probability(확률), aml_patterns(탐지 패턴) 파라미터를 "
        "통해 보고서 섹션을 자동 생성한다.",
        [
            ("required", "summary", "string",
             "분석 결과 요약 및 혐의 판단 사유 (의심 사유·거래 패턴·수치 등 구체적으로 기술)"),
            ("optional", "fraud_type", "string",
             "이상거래유형: 자금세탁 / 대포통장 / 보이스피싱 / 불법도박 / 유사수신 / 신규거래처 / 기타"),
            ("optional", "transactions", "array",
             "query_transactions 결과에서 가져온 관련 거래 레코드 목록"),
            ("optional", "fraud_probability", "number",
             "predict_fraud 예측 이상거래 확률 (0.0~1.0). 의심 강도(1~5) 산출에 사용."),
            ("optional", "aml_patterns", "array",
             "detect_aml_patterns로 탐지된 패턴 목록 (예: ['순환거래', '레이어링'])"),
        ],
    ),

    # ===================================================================
    # Sub domain 3: CTR·위험평가·모니터링
    # ===================================================================

    # --- detect_ctr_candidates (① 정답 도구) ---
    (
        "CTR·위험평가·모니터링",
        "CTR 대상 고액거래 및 분할거래 탐지 (① 정답 도구)",
        "detect_ctr_candidates",
        "1,000만원 이상 고액현금거래(CTR) 보고 대상을 탐지하거나, "
        "동일계좌 동일일 합산이 기준금액 이상이면서 단건이 기준 미만인 분할거래(structuring) 패턴을 탐지한다.",
        [
            ("required", "mode", "string",
             "탐지 모드: high_value(고액거래) / structuring(분할거래)"),
            ("optional", "date_from", "integer", "시작일 (YYYYMMDD)"),
            ("optional", "date_to", "integer", "종료일 (YYYYMMDD)"),
            ("optional", "threshold", "integer", "기준금액 (기본값 10000000)"),
            ("optional", "limit", "integer", "반환 건수 (기본값 20)"),
        ],
    ),

    # --- score_account_risk (① 정답 도구) ---
    (
        "CTR·위험평가·모니터링",
        "계좌 행위 기반 위험도 종합 평가 (① 정답 도구)",
        "score_account_risk",
        "5개 행위 지표(심야거래비율, 금액이상도, 거래상대다양성, 거래속도변화, 이상거래이력)를 "
        "가중 합산하여 계좌의 종합 위험도를 0~100점으로 산출한다.",
        [
            ("required", "account_id", "integer",
             "평가할 계좌 번호 (출금계좌일련번호). 예: 1234567890"),
        ],
    ),

    # --- detect_monitoring_alerts (① 정답 도구) ---
    (
        "CTR·위험평가·모니터링",
        "규칙 기반 거래 모니터링 알림 탐지 (① 정답 도구)",
        "detect_monitoring_alerts",
        "5개 모니터링 규칙(R001 심야대량, R002 동일일다건, R003 정액패턴, "
        "R004 기관집중, R005 패턴급변)을 실행하여 알림을 생성한다.",
        [
            ("required", "rule_id", "string",
             "규칙 ID: all / R001 / R002 / R003 / R004 / R005"),
            ("required", "date_from", "integer", "시작일 (YYYYMMDD)"),
            ("required", "date_to", "integer", "종료일 (YYYYMMDD)"),
            ("optional", "account_id", "integer", "특정 계좌 한정 (선택)"),
            ("optional", "limit", "integer", "반환 건수 (기본값 20)"),
        ],
    ),

    # ===================================================================
    # Sub domain 4: 자금 흐름·추세·채널 분석
    # ===================================================================

    # --- detect_dormant_reactivation (① 정답 도구) ---
    (
        "자금 흐름·추세·채널 분석",
        "장기 휴면 계좌 재활성화 탐지 (① 정답 도구)",
        "detect_dormant_reactivation",
        "일정 기간(기본 180일) 이상 거래가 없다가 재활성화된 계좌를 탐지한다. "
        "자금세탁 목적의 휴면계좌 활용 패턴 식별에 사용한다.",
        [
            ("optional", "dormant_days", "integer", "휴면 기준 일수 (기본값 180)"),
            ("optional", "min_reactivation_amount", "integer",
             "최소 재활성화 거래금액 (기본값 5000000)"),
            ("optional", "limit", "integer", "반환 건수 (기본값 100)"),
        ],
    ),

    # --- detect_smurfing_network (① 정답 도구) ---
    (
        "자금 흐름·추세·채널 분석",
        "스머핑(자금 수집/분산) 네트워크 탐지 (① 정답 도구)",
        "detect_smurfing_network",
        "다수 계좌에서 하나의 계좌로 자금을 모으는 수집(inbound) 패턴이나, "
        "하나의 계좌에서 다수 계좌로 분산(outbound)하는 패턴을 탐지한다.",
        [
            ("optional", "account_id", "integer", "특정 계좌 한정 (선택)"),
            ("required", "direction", "string",
             "분석 방향: inbound(자금 수집) / outbound(자금 분산)"),
            ("optional", "min_counterparts", "integer", "최소 거래상대 수 (기본값 5)"),
            ("optional", "date_from", "integer", "시작일 (YYYYMMDD)"),
            ("optional", "date_to", "integer", "종료일 (YYYYMMDD)"),
            ("optional", "limit", "integer", "반환 건수 (기본값 50)"),
        ],
    ),

    # --- get_trend_analysis (① 정답 도구) ---
    (
        "자금 흐름·추세·채널 분석",
        "시계열 추세 분석 (① 정답 도구)",
        "get_trend_analysis",
        "월별 또는 분기별 거래 추세를 분석한다. 거래건수, 거래금액, 이상거래건수, "
        "이상거래비율 중 선택한 지표의 시계열 데이터를 반환한다.",
        [
            ("optional", "unit", "string",
             "집계 단위: monthly(월별, 기본값) / quarterly(분기별)"),
            ("optional", "metric", "string",
             "지표: count(거래건수, 기본값) / amount(거래금액) / fraud_count(이상거래건수) / fraud_rate(이상거래비율)"),
            ("optional", "date_from", "integer", "시작일 (YYYYMMDD)"),
            ("optional", "date_to", "integer", "종료일 (YYYYMMDD)"),
        ],
    ),

    # --- analyze_channel_risk (① 정답 도구) ---
    (
        "자금 흐름·추세·채널 분석",
        "채널(매체) 위험도 분석 (① 정답 도구)",
        "analyze_channel_risk",
        "매체구분(채널)별 거래 통계와 이상거래 비율, 채널-시간대 교차 분석 결과를 반환한다. "
        "특정 채널의 위험도를 파악하는 데 사용한다.",
        [
            ("optional", "date_from", "integer", "시작일 (YYYYMMDD)"),
            ("optional", "date_to", "integer", "종료일 (YYYYMMDD)"),
        ],
    ),

    # --- get_receiving_account_profile (① 정답 도구) ---
    (
        "자금 흐름·추세·채널 분석",
        "입금 계좌 프로파일 조회 (① 정답 도구)",
        "get_receiving_account_profile",
        "입금계좌 기준으로 거래 프로파일을 조회한다. "
        "총 거래 건수·금액, 이상거래 비율, 상위 출금 계좌, 상위 출금 금융회사, 시간대별 분포를 반환한다. "
        "get_account_profile(출금계좌 기준)의 반대 방향 조회이다.",
        [
            ("required", "account_id", "integer",
             "조회할 입금계좌 번호 (입금계좌일련번호 기준). 예: 1234567890"),
        ],
    ),

    # --- analyze_cross_institution_flow (① 정답 도구) ---
    (
        "자금 흐름·추세·채널 분석",
        "기관간 자금 흐름 분석 (① 정답 도구)",
        "analyze_cross_institution_flow",
        "출금금융회사 → 입금금융회사 간 자금 흐름을 분석한다. "
        "기관 쌍별 거래건수, 이상거래건수, 이상거래비율을 반환하여 "
        "기관간 이상 자금 흐름 패턴을 식별한다.",
        [
            ("optional", "date_from", "integer", "시작일 (YYYYMMDD)"),
            ("optional", "date_to", "integer", "종료일 (YYYYMMDD)"),
            ("optional", "min_transactions", "integer", "최소 거래건수 필터 (기본값 10)"),
            ("optional", "limit", "integer", "반환 건수 (기본값 50)"),
        ],
    ),
]

# ---------------------------------------------------------------------------
# ② 평가 시나리오 데이터
# ---------------------------------------------------------------------------
# 각 항목: (no, sub_domain, task, 시나리오설정, 사용자질문, func_name,
#           [(param_name, value), ...], 근거, 사고과정)

_SCENARIOS: list[tuple] = [

    # ── get_statistics ──────────────────────────────────────────────────
    (
        1, "이상거래 통계 및 현황 조회",
        "전체 이상거래 요약 통계 조회 (① 정답 도구)",
        "이상거래 모니터링 담당 직원이 매일 아침 업무 시작 시 전체 이상거래 현황을 빠르게 파악해야 하는 상황",
        "오늘 업무 시작 전에 전체 이상거래 현황을 한 번에 확인하고 싶어. 유형별 분포도 보여줘.",
        "get_statistics", [],
        "전체 이상거래 건수·비율·유형별 분포 등 요약 통계를 조회하는 것이 목적임. "
        "파라미터 없이 호출하면 전체 기간 전체 유형 집계를 반환하므로 get_statistics가 적합. "
        "query_transactions는 row 단위 반환으로 요약 통계를 바로 얻을 수 없고, "
        "get_fraud_type_summary는 유형 코드를 지정해야 하므로 전체 요약 조회에 부적합.",
        "[1. 질문 주제 파악]\n• 사용자는 AML 이상거래 모니터링 담당 직원이며, 전체 이상거래 요약 현황 조회가 주제임.\n\n"
        "[2. 질문 목적 분석]\n• 업무 시작 전 전체 거래 건수, 이상거래 건수·비율, 유형별 분포를 한 번에 파악하고자 함.\n\n"
        "[3. 필요 기능 도출]\n• 파라미터 없이 전체 집계 요약을 반환하는 get_statistics가 최적. "
        "query_transactions는 row 단위, get_fraud_type_summary는 유형 코드 지정이 필요하여 부적합.\n\n"
        "[4. 필수 파라미터 추출]\n• 파라미터 없음 → get_statistics()",
    ),
    (
        2, "이상거래 통계 및 현황 조회",
        "전체 이상거래 요약 통계 조회 (① 정답 도구)",
        "팀장이 주간 보고를 위해 HOFINET 전체 거래 규모와 이상거래 비율을 확인해야 하는 상황",
        "HOFINET 데이터에서 전체 거래 건수랑 이상거래 비율이 어떻게 돼?",
        "get_statistics", [],
        "전체 거래 건수와 이상거래 비율을 조회하는 것이 목적이므로 get_statistics가 적합. "
        "특정 유형이나 계좌를 지정하지 않았으므로 get_fraud_type_summary나 "
        "get_account_profile은 사용할 수 없음. "
        "query_transactions로도 집계 가능하나 SQL을 직접 작성해야 하므로 비효율적.",
        "[1. 질문 주제 파악]\n• 사용자는 팀장 보고를 준비하는 AML 담당 직원이며, 전체 거래·이상거래 통계 조회가 주제임.\n\n"
        "[2. 질문 목적 분석]\n• 전체 거래 건수와 이상거래 비율을 파악하여 주간 보고 자료를 준비하고자 함.\n\n"
        "[3. 필요 기능 도출]\n• 전체 요약 통계를 반환하는 get_statistics가 적합. "
        "특정 유형 코드나 계좌 번호를 모르는 상황에서 다른 도구를 사용하면 추가 정보가 필요함.\n\n"
        "[4. 필수 파라미터 추출]\n• 파라미터 없음 → get_statistics()",
    ),

    # ── get_fraud_type_summary ───────────────────────────────────────────
    (
        3, "이상거래 통계 및 현황 조회",
        "이상거래 유형별 상세 현황 조회 (① 정답 도구)",
        "수사 협조 요청으로 보이스피싱 의심거래의 전체 건수·금액·관련 금융회사를 파악해야 하는 상황",
        "보이스피싱 의심거래 건수와 총 금액이 얼마야? 어느 금융회사에서 많이 발생했는지도 알려줘.",
        "get_fraud_type_summary", [("fraud_type", "4")],
        "보이스피싱 유형(코드 4)의 건수·금액·상위 금융회사 집계가 목적. "
        "get_fraud_type_summary는 특정 유형의 집계 통계와 상위 금융회사를 반환하므로 적합. "
        "get_statistics는 전체 유형 요약만 반환하여 유형별 상세 분석 불가. "
        "query_transactions는 row 단위 반환으로 집계 결과를 바로 얻을 수 없음.",
        "[1. 질문 주제 파악]\n• 사용자는 수사 협조 담당 AML 직원이며, 보이스피싱 이상거래 통계 조회가 주제임.\n\n"
        "[2. 질문 목적 분석]\n• 보이스피싱 의심거래의 전체 건수, 총 금액, 상위 관련 금융회사를 파악하고자 함.\n\n"
        "[3. 필요 기능 도출]\n• 특정 유형 코드로 집계 통계와 상위 금융회사를 반환하는 get_fraud_type_summary가 적합. "
        "보이스피싱 = 코드 4. get_statistics는 전체 요약만, query_transactions는 row 단위로 부적합.\n\n"
        "[4. 필수 파라미터 추출]\n• 보이스피싱 → fraud_type=4",
    ),
    (
        4, "이상거래 통계 및 현황 조회",
        "이상거래 유형별 상세 현황 조회 (① 정답 도구)",
        "특정 금융회사에서 발생한 자금세탁 의심거래 현황을 집중 점검해야 하는 상황",
        "금융회사 번호 101번에서 발생한 자금세탁(유형1) 이상거래 현황을 알려줘.",
        "get_fraud_type_summary", [("fraud_type", "1"), ("bank_id", "101")],
        "자금세탁(코드 1) + 금융회사 101번 필터 집계가 목적. "
        "get_fraud_type_summary는 fraud_type 필수 파라미터와 bank_id 선택 파라미터를 모두 지원하여 적합. "
        "get_account_profile은 계좌 단위이므로 금융회사 전체 집계 불가. "
        "query_transactions는 집계용 도구가 아니므로 부적합.",
        "[1. 질문 주제 파악]\n• 사용자는 금융회사 집중 모니터링 담당 직원이며, 특정 금융회사의 자금세탁 이상거래 현황 조회가 주제임.\n\n"
        "[2. 질문 목적 분석]\n• 금융회사 101번에서 발생한 자금세탁(유형1) 건수·금액을 파악하여 점검 자료를 준비하고자 함.\n\n"
        "[3. 필요 기능 도출]\n• fraud_type(필수)과 bank_id(선택) 필터를 모두 지원하는 get_fraud_type_summary가 최적. "
        "자금세탁=코드 1, 금융회사=bank_id 101.\n\n"
        "[4. 필수 파라미터 추출]\n• 자금세탁 → fraud_type=1\n• 금융회사 번호 101 → bank_id=101",
    ),

    # ── get_account_profile ─────────────────────────────────────────────
    (
        5, "이상거래 통계 및 현황 조회",
        "특정 계좌의 거래 통계 프로파일 조회 (② 유사 기능 도구)",
        "민원 처리 중 특정 계좌의 이상거래 이력과 거래 패턴을 신속하게 파악해야 하는 상황",
        "계좌 번호 9876543210의 전체 거래 통계와 이상거래 비율을 알려줘.",
        "get_account_profile", [("account_id", "9876543210")],
        "단일 계좌의 거래 통계(건수·금액·이상거래비율·상위 상대 계좌)를 조회하는 것이 목적. "
        "get_account_profile은 account_id 하나로 계좌 단위 집계를 반환하므로 적합. "
        "get_statistics는 전체 현황만 반환하므로 특정 계좌 분석 불가. "
        "analyze_network는 네트워크 연결 구조 분석이 목적으로 통계 프로파일 반환 불가.",
        "[1. 질문 주제 파악]\n• 사용자는 AML 민원 처리 담당 직원이며, 특정 계좌의 거래 통계 조회가 주제임.\n\n"
        "[2. 질문 목적 분석]\n• 계좌 9876543210의 전체 거래 건수·금액·이상거래 비율을 파악하고자 함.\n\n"
        "[3. 필요 기능 도출]\n• account_id를 입력하면 계좌 단위 집계 통계를 반환하는 get_account_profile이 적합. "
        "get_statistics는 전체 집계, analyze_network는 연결 구조 분석으로 부적합.\n\n"
        "[4. 필수 파라미터 추출]\n• 계좌 번호 9876543210 → account_id=9876543210",
    ),
    (
        6, "이상거래 통계 및 현황 조회",
        "특정 계좌의 거래 통계 프로파일 조회 (② 유사 기능 도구)",
        "내부 감사팀이 의심 계좌의 주요 거래 상대와 선호 거래 시간대를 파악해야 하는 상황",
        "계좌 1111222233에서 가장 많이 거래한 상대 계좌가 어디야? 주로 몇 시에 거래했어?",
        "get_account_profile", [("account_id", "1111222233")],
        "상위 거래 상대 계좌와 주요 거래 시간대를 조회하는 것이 목적. "
        "get_account_profile의 반환값에 top_counterparts와 top_hours가 포함되어 있으므로 적합. "
        "query_transactions로도 SQL로 집계할 수 있으나 도구 호출 목적이 명확하지 않고 복잡. "
        "analyze_network는 연결 계좌 수·이상거래 여부만 반환하며 거래 시간대는 미포함.",
        "[1. 질문 주제 파악]\n• 사용자는 내부 감사 담당 직원이며, 특정 계좌의 거래 상대 및 시간대 패턴 조회가 주제임.\n\n"
        "[2. 질문 목적 분석]\n• 계좌 1111222233의 상위 거래 상대 계좌 목록과 주요 거래 시간대를 파악하고자 함.\n\n"
        "[3. 필요 기능 도출]\n• get_account_profile은 top_counterparts(상위 거래 상대)와 top_hours(주요 시간대)를 반환하므로 적합. "
        "analyze_network는 거래 시간대 정보가 없고, query_transactions는 집계 도구가 아님.\n\n"
        "[4. 필수 파라미터 추출]\n• 계좌 번호 1111222233 → account_id=1111222233",
    ),

    # ── query_transactions ──────────────────────────────────────────────
    (
        7, "이상거래 통계 및 현황 조회",
        "SQL 기반 거래 데이터 원본 조회 (③ 인접·미끼 도구)",
        "특정 날짜에 발생한 고액 이상거래 레코드를 직접 확인해야 하는 상황",
        "20221015일에 이상거래가 발생한 건 중 거래금액이 5000만원 이상인 건만 조회해줘.",
        "query_transactions",
        [("sql", "SELECT * FROM hofinet WHERE 거래일자=20221015 AND 이상거래여부=1 AND 거래금액>=50000000")],
        "특정 날짜와 금액 조건으로 이상거래 원본 레코드를 직접 조회하는 것이 목적. "
        "query_transactions는 조건부 SQL을 실행하여 row 단위 결과를 반환하므로 적합. "
        "get_statistics는 집계 요약만 반환하여 특정 날짜·금액 필터 불가. "
        "get_fraud_type_summary는 유형별 집계로 개별 레코드 조회 불가.",
        "[1. 질문 주제 파악]\n• 사용자는 AML 조사 담당 직원이며, 특정 날짜·금액 조건의 이상거래 레코드 조회가 주제임.\n\n"
        "[2. 질문 목적 분석]\n• 20221015일에 발생한 거래금액 5000만원 이상의 이상거래 레코드를 직접 확인하고자 함.\n\n"
        "[3. 필요 기능 도출]\n• 날짜·금액 조건을 WHERE절로 명시한 SELECT SQL을 실행할 수 있는 query_transactions가 적합. "
        "get_statistics와 get_fraud_type_summary는 집계 통계 반환으로 개별 레코드 조회 불가.\n\n"
        "[4. 필수 파라미터 추출]\n• sql=\"SELECT * FROM hofinet WHERE 거래일자=20221015 AND 이상거래여부=1 AND 거래금액>=50000000\"",
    ),
    (
        8, "이상거래 통계 및 현황 조회",
        "SQL 기반 거래 데이터 원본 조회 (③ 인접·미끼 도구)",
        "매체구분별 이상거래 건수 분포를 직접 집계하여 보고서 초안을 작성해야 하는 상황",
        "매체구분(채널)별로 이상거래가 몇 건이나 발생했는지 집계해줘.",
        "query_transactions",
        [("sql", "SELECT 매체구분, COUNT(*) AS 건수 FROM hofinet WHERE 이상거래여부=1 GROUP BY 매체구분 ORDER BY 건수 DESC")],
        "매체구분별 이상거래 건수를 그룹핑하여 집계하는 것이 목적. "
        "query_transactions는 GROUP BY SQL을 직접 실행할 수 있으므로 적합. "
        "get_statistics는 유형별 분포만 반환하며 매체구분 기준 집계는 포함되지 않음. "
        "get_fraud_type_summary는 유형 코드를 지정해야 하므로 매체구분 집계에 부적합.",
        "[1. 질문 주제 파악]\n• 사용자는 AML 보고서 작성 담당 직원이며, 매체구분별 이상거래 건수 집계가 주제임.\n\n"
        "[2. 질문 목적 분석]\n• 채널(매체구분)별 이상거래 건수 분포를 집계하여 보고서에 활용하고자 함.\n\n"
        "[3. 필요 기능 도출]\n• GROUP BY 매체구분 SQL을 실행할 수 있는 query_transactions가 적합. "
        "get_statistics는 유형별 분포만 포함, get_fraud_type_summary는 유형 지정 필요로 부적합.\n\n"
        "[4. 필수 파라미터 추출]\n• sql=\"SELECT 매체구분, COUNT(*) AS 건수 FROM hofinet WHERE 이상거래여부=1 GROUP BY 매체구분 ORDER BY 건수 DESC\"",
    ),

    # ── analyze_network ─────────────────────────────────────────────────
    (
        9, "AML 탐지·분석 및 보고서 작성",
        "계좌 거래 네트워크 구조 분석 (① 정답 도구)",
        "자금세탁 혐의 계좌가 보고되어 해당 계좌의 거래 연결망 규모와 이상거래 관련 여부를 신속하게 파악해야 하는 상황",
        "계좌 5551234567과 연결된 계좌가 몇 개나 되고, 그중 이상거래 관련 계좌가 있어?",
        "analyze_network", [("account_id", "5551234567")],
        "특정 계좌의 거래 네트워크 연결 계좌 수와 이상거래 관련 여부를 파악하는 것이 목적. "
        "analyze_network는 account_id로 연결 계좌 수·거래 횟수·이상거래 연관 여부를 반환하므로 적합. "
        "detect_aml_patterns는 특정 AML 패턴(ring/layering 등) 탐지 목적으로 네트워크 규모 파악에는 부적합. "
        "get_account_profile은 거래 통계 프로파일로 연결망 구조는 포함하지 않음.",
        "[1. 질문 주제 파악]\n• 사용자는 AML 혐의 계좌 분석 담당 직원이며, 계좌 거래 네트워크 규모 파악이 주제임.\n\n"
        "[2. 질문 목적 분석]\n• 계좌 5551234567에 연결된 계좌 수와 이상거래 연관 여부를 확인하고자 함.\n\n"
        "[3. 필요 기능 도출]\n• 연결 계좌 수·이상거래 연관 여부를 반환하는 analyze_network가 적합. "
        "detect_aml_patterns는 패턴 유형 지정이 필요하고 연결망 규모 반환이 주목적이 아님.\n\n"
        "[4. 필수 파라미터 추출]\n• 계좌 번호 5551234567 → account_id=5551234567",
    ),
    (
        10, "AML 탐지·분석 및 보고서 작성",
        "계좌 거래 네트워크 구조 분석 (① 정답 도구)",
        "심층 자금 흐름 분석을 위해 특정 계좌의 2단계 이상 거래 연결망을 확인해야 하는 상황",
        "계좌 3339876543에 대해 2-hop 거래 네트워크를 분석해줘.",
        "analyze_network", [("account_id", "3339876543"), ("hops", "2")],
        "2-hop 심층 탐색으로 간접 연결 계좌까지 포함한 네트워크 분석이 목적. "
        "analyze_network는 hops 파라미터로 탐색 깊이를 지정할 수 있으므로 적합. "
        "detect_aml_patterns의 shortest_path는 두 계좌 간 경로 탐색이 목적으로 "
        "단일 계좌의 N-hop 네트워크 분석과는 다름.",
        "[1. 질문 주제 파악]\n• 사용자는 AML 심층 분석 담당 직원이며, 2-hop 거래 네트워크 분석이 주제임.\n\n"
        "[2. 질문 목적 분석]\n• 계좌 3339876543의 직접 거래 계좌(1-hop)뿐 아니라 그 계좌들과 연결된 계좌(2-hop)까지 확인하고자 함.\n\n"
        "[3. 필요 기능 도출]\n• hops 파라미터를 지원하는 analyze_network가 적합. "
        "detect_aml_patterns는 패턴 탐지 목적으로 N-hop 네트워크 구조 반환은 하지 않음.\n\n"
        "[4. 필수 파라미터 추출]\n• 계좌 번호 3339876543 → account_id=3339876543\n• 탐색 깊이 2-hop → hops=2",
    ),

    # ── detect_aml_patterns ─────────────────────────────────────────────
    (
        11, "AML 탐지·분석 및 보고서 작성",
        "AML 특화 패턴 탐지 (② 유사 기능 도구)",
        "자금세탁 혐의 계좌에 대해 순환거래 패턴이 있는지 그래프 DB로 확인해야 하는 상황",
        "Memgraph에서 순환거래 패턴을 탐지해줘.",
        "detect_aml_patterns", [("pattern_type", "ring")],
        "순환거래(ring) 패턴을 Memgraph 그래프 DB로 탐지하는 것이 목적. "
        "detect_aml_patterns의 pattern_type=\"ring\"이 정확히 이 기능을 수행함. "
        "analyze_network는 연결 구조·규모 분석이지 순환거래 패턴 탐지가 아님. "
        "predict_fraud는 개별 거래 ML 예측으로 패턴 탐지와 무관.",
        "[1. 질문 주제 파악]\n• 사용자는 AML 그래프 분석 담당 직원이며, 순환거래 패턴 탐지가 주제임.\n\n"
        "[2. 질문 목적 분석]\n• Memgraph 그래프 DB를 활용하여 순환거래(ring) 패턴이 존재하는지 탐지하고자 함.\n\n"
        "[3. 필요 기능 도출]\n• detect_aml_patterns의 pattern_type=\"ring\"이 Memgraph 기반 순환거래 탐지를 수행. "
        "analyze_network는 연결망 규모 분석이고, predict_fraud는 ML 예측으로 부적합.\n\n"
        "[4. 필수 파라미터 추출]\n• 순환거래 탐지 → pattern_type=\"ring\"",
    ),
    (
        12, "AML 탐지·분석 및 보고서 작성",
        "AML 특화 패턴 탐지 (② 유사 기능 도구)",
        "의심 계좌의 위험도를 정량 평가하여 STR 작성 여부를 결정해야 하는 상황",
        "계좌 7778889990의 AML 위험도 점수를 산출해줘.",
        "detect_aml_patterns", [("pattern_type", "risk_score"), ("account_id", "7778889990")],
        "특정 계좌의 AML 위험도 점수(risk_score) 산출이 목적. "
        "detect_aml_patterns의 pattern_type=\"risk_score\"와 account_id 조합이 정확히 이 기능 수행. "
        "analyze_network는 네트워크 구조만 반환하며 정량 위험도 점수는 포함하지 않음. "
        "predict_fraud는 개별 거래 단위 확률 예측으로 계좌 위험도와 다름.",
        "[1. 질문 주제 파악]\n• 사용자는 STR 작성 여부 판단을 담당하는 AML 직원이며, 계좌 위험도 점수 산출이 주제임.\n\n"
        "[2. 질문 목적 분석]\n• 계좌 7778889990의 AML 위험도를 정량적으로 평가하여 STR 작성 필요 여부를 판단하고자 함.\n\n"
        "[3. 필요 기능 도출]\n• detect_aml_patterns(pattern_type=\"risk_score\", account_id=...)가 계좌 위험도 점수를 산출. "
        "analyze_network는 위험도 점수 없음, predict_fraud는 거래 단위 예측으로 부적합.\n\n"
        "[4. 필수 파라미터 추출]\n• 위험도 산출 → pattern_type=\"risk_score\"\n• 계좌 번호 7778889990 → account_id=7778889990",
    ),

    # ── predict_fraud ───────────────────────────────────────────────────
    (
        13, "AML 탐지·분석 및 보고서 작성",
        "XGBoost 모델 기반 이상거래 확률 예측 (③ 인접·미끼 도구)",
        "신규 입력된 단건 거래가 이상거래인지 즉시 판별해야 하는 상황",
        "자정에 핀테크 앱으로 100만원 이체한 거래가 이상거래일 확률이 얼마야? "
        "출금기관 201, 입금기관 305, 자금구분 1, 매체구분 5야.",
        "predict_fraud",
        [("거래시간대", "0"), ("출금금융회사일련번호", "201"), ("입금금융회사일련번호", "305"),
         ("자금구분", "1"), ("매체구분", "5"), ("거래금액", "1000000")],
        "단건 거래 속성 6개를 입력하여 XGBoost 모델로 이상거래 확률을 예측하는 것이 목적. "
        "predict_fraud는 거래 속성을 입력받아 0~1 확률을 반환하므로 적합. "
        "analyze_network는 계좌 네트워크 분석으로 단건 확률 예측과 무관. "
        "detect_aml_patterns는 그래프 패턴 탐지로 단건 ML 예측과 다름.",
        "[1. 질문 주제 파악]\n• 사용자는 실시간 이상거래 모니터링 담당 직원이며, 단건 거래의 이상거래 확률 예측이 주제임.\n\n"
        "[2. 질문 목적 분석]\n• 자정(시간대 0), 핀테크 앱(매체구분 5), 100만원 거래가 이상거래일 확률을 즉시 판별하고자 함.\n\n"
        "[3. 필요 기능 도출]\n• 거래 속성 6개를 입력받아 XGBoost로 0~1 확률을 반환하는 predict_fraud가 적합. "
        "analyze_network와 detect_aml_patterns는 계좌·그래프 분석 도구로 단건 ML 예측과 무관.\n\n"
        "[4. 필수 파라미터 추출]\n• 자정 → 거래시간대=0\n• 출금기관=201, 입금기관=305, 자금구분=1, 매체구분=5, 거래금액=1000000",
    ),
    (
        14, "AML 탐지·분석 및 보고서 작성",
        "XGBoost 모델 기반 이상거래 확률 예측 (③ 인접·미끼 도구)",
        "STR 작성을 위해 의심 거래의 AI 예측 확률을 먼저 확인해야 하는 상황",
        "오전 9시에 인터넷뱅킹으로 출금기관 88, 입금기관 77, 자금구분 4, 500만원을 이체한 거래의 이상거래 확률을 예측해줘.",
        "predict_fraud",
        [("거래시간대", "9"), ("출금금융회사일련번호", "88"), ("입금금융회사일련번호", "77"),
         ("자금구분", "4"), ("매체구분", "2"), ("거래금액", "5000000")],
        "STR 작성 전 단계로 특정 거래의 이상거래 확률을 예측하는 것이 목적. "
        "predict_fraud는 6개 거래 속성을 입력받아 ML 확률을 반환하므로 적합. "
        "인터넷뱅킹=매체구분 2, 오전 9시=거래시간대 9. "
        "generate_str는 STR 보고서 작성 도구로 확률 예측 기능이 없음.",
        "[1. 질문 주제 파악]\n• 사용자는 STR 작성 전 AI 예측을 수행하는 AML 직원이며, 단건 거래 확률 예측이 주제임.\n\n"
        "[2. 질문 목적 분석]\n• 오전 9시 인터넷뱅킹 500만원 이체 거래가 이상거래일 확률을 ML 모델로 예측하고자 함.\n\n"
        "[3. 필요 기능 도출]\n• predict_fraud가 6개 속성 입력 → XGBoost 확률 반환. "
        "인터넷뱅킹=매체구분 2, 오전 9시=거래시간대 9 매핑. generate_str는 확률 예측 기능 없음.\n\n"
        "[4. 필수 파라미터 추출]\n• 오전 9시 → 거래시간대=9\n• 인터넷뱅킹 → 매체구분=2\n• 출금기관=88, 입금기관=77, 자금구분=4, 거래금액=5000000",
    ),

    # ── generate_str ────────────────────────────────────────────────────
    (
        15, "AML 탐지·분석 및 보고서 작성",
        "의심거래보고서(STR) 공식 양식 작성 (① 정답 도구)",
        "분석 완료 후 자금세탁 의심 계좌에 대한 STR을 공식 양식(I~VII 섹션)에 맞게 작성 제출해야 하는 상황",
        "보이스피싱 의심 거래에 대해 STR 보고서를 작성해줘. "
        "혐의 사유는 '자정 시간대 반복 소액 분산 입금 후 즉시 전액 출금'이야.",
        "generate_str",
        [("summary", "자정 시간대 반복 소액 분산 입금 후 즉시 전액 출금"), ("fraud_type", "보이스피싱")],
        "분석 결과를 바탕으로 STR 공식 양식을 작성하는 것이 목적. "
        "generate_str는 summary(필수)와 fraud_type(선택)을 입력받아 I~VII 섹션 구조의 STR을 생성하므로 적합. "
        "predict_fraud는 확률 예측 도구로 STR 작성 기능이 없음. "
        "query_transactions는 데이터 조회 도구로 보고서 작성과 무관.",
        "[1. 질문 주제 파악]\n• 사용자는 STR 작성 담당 AML 직원이며, 의심거래보고서 공식 양식 작성이 주제임.\n\n"
        "[2. 질문 목적 분석]\n• 보이스피싱 의심 거래에 대해 혐의 사유를 기반으로 I~VII 섹션 구조의 STR을 작성하고자 함.\n\n"
        "[3. 필요 기능 도출]\n• generate_str(summary=..., fraud_type=\"보이스피싱\")이 STR 공식 양식을 생성. "
        "predict_fraud는 확률 예측, query_transactions는 데이터 조회로 보고서 작성 기능 없음.\n\n"
        "[4. 필수 파라미터 추출]\n• 혐의 사유: '자정 시간대 반복 소액 분산 입금 후 즉시 전액 출금' → summary=...\n• 보이스피싱 → fraud_type=\"보이스피싱\"",
    ),

    # ── compare_periods ─────────────────────────────────────────────────
    (
        16, "이상거래 통계 및 현황 조회",
        "두 기간 이상거래 통계 비교 (① 정답 도구)",
        "분기 보고를 위해 2024년 Q4와 Q3의 이상거래 건수·금액·비율 변화를 분석해야 하는 상황",
        "2024년 4분기(10~12월)와 3분기(7~9월)의 이상거래 통계를 비교해줘. 건수와 비율 변화가 궁금해.",
        "compare_periods",
        [("period1_start", "20240701"), ("period1_end", "20240930"),
         ("period2_start", "20241001"), ("period2_end", "20241231")],
        "두 분기의 이상거래 건수·비율·평균 금액 변화를 비교하는 것이 목적. "
        "compare_periods는 두 날짜 범위를 입력받아 각 기간의 집계와 delta를 반환하므로 적합. "
        "get_statistics는 전체 기간 집계만 반환하여 기간 비교 불가. "
        "query_transactions로도 두 번 집계 가능하나 delta 계산 없이 row 단위 반환으로 비효율적.",
        "[1. 질문 주제 파악]\n• 사용자는 분기 보고 담당 AML 직원이며, 2분기 간 이상거래 통계 비교가 주제임.\n\n"
        "[2. 질문 목적 분석]\n• Q3(7~9월) vs Q4(10~12월)의 이상거래 건수·비율 변화를 파악하여 보고 자료를 준비하고자 함.\n\n"
        "[3. 필요 기능 도출]\n• compare_periods는 period1/period2 날짜 범위를 입력받아 집계 통계와 delta를 반환하므로 적합. "
        "get_statistics는 기간 필터 없음, query_transactions는 집계·비교 기능 없음.\n\n"
        "[4. 필수 파라미터 추출]\n• Q3: period1_start=20240701, period1_end=20240930\n• Q4: period2_start=20241001, period2_end=20241231",
    ),
    (
        17, "이상거래 통계 및 현황 조회",
        "두 기간 이상거래 통계 비교 (① 정답 도구)",
        "연간 이상거래 추이 파악을 위해 2023년과 2022년 전체 연도의 이상거래 현황을 비교해야 하는 상황",
        "2023년과 2022년 이상거래 통계를 연도별로 비교해줘. 어느 해에 더 많이 발생했어?",
        "compare_periods",
        [("period1_start", "20220101"), ("period1_end", "20221231"),
         ("period2_start", "20230101"), ("period2_end", "20231231")],
        "2022년과 2023년 전체의 이상거래 건수·비율을 연간 단위로 비교하는 것이 목적. "
        "compare_periods는 YYYYMMDD 범위로 연간 집계 비교와 delta를 반환하므로 적합. "
        "get_statistics는 전체 기간 통합 집계만 반환하여 연도별 분리 비교 불가. "
        "get_fraud_type_summary는 특정 유형 지정이 필요하여 전체 이상거래 연간 비교에 부적합.",
        "[1. 질문 주제 파악]\n• 사용자는 연간 추이 분석 담당 AML 직원이며, 2022년과 2023년 이상거래 연간 비교가 주제임.\n\n"
        "[2. 질문 목적 분석]\n• 두 연도의 이상거래 건수 및 비율을 비교하여 추이 변화를 파악하고자 함.\n\n"
        "[3. 필요 기능 도출]\n• compare_periods에 1월1일~12월31일 범위를 각 연도별로 입력하면 연간 집계 비교가 가능. "
        "get_statistics는 전체 기간 요약만, get_fraud_type_summary는 유형별 집계로 부적합.\n\n"
        "[4. 필수 파라미터 추출]\n• 2022년: period1_start=20220101, period1_end=20221231\n• 2023년: period2_start=20230101, period2_end=20231231",
    ),

    # ── get_institution_report ──────────────────────────────────────────
    (
        18, "이상거래 통계 및 현황 조회",
        "금융회사 종합 현황 보고 (① 정답 도구)",
        "감독원 현장검사 준비를 위해 특정 금융회사의 이상거래 현황·거래 규모·주요 거래 상대 기관을 파악해야 하는 상황",
        "금융회사 번호 134의 이상거래 비율, 거래 규모, 주요 거래 상대 기관을 종합적으로 알려줘.",
        "get_institution_report", [("bank_id", "134")],
        "단일 금융회사의 종합 현황(거래 건수·이상거래 비율·거래 상대 기관·유형 분포)을 조회하는 것이 목적. "
        "get_institution_report는 bank_id 하나로 금융회사 단위 종합 집계를 반환하므로 적합. "
        "get_fraud_type_summary는 특정 이상거래 유형 집계가 목적이며 기관 종합 현황 반환 불가. "
        "query_transactions는 row 단위 반환으로 기관 단위 종합 집계 불가.",
        "[1. 질문 주제 파악]\n• 사용자는 현장검사 준비 담당 AML 직원이며, 특정 금융회사 종합 현황 조회가 주제임.\n\n"
        "[2. 질문 목적 분석]\n• 금융회사 134번의 이상거래 비율, 거래 규모, 상위 거래 상대 기관을 한 번에 파악하고자 함.\n\n"
        "[3. 필요 기능 도출]\n• get_institution_report(bank_id=134)는 거래 건수·이상거래 비율·상위 상대 기관·유형 분포를 반환. "
        "get_fraud_type_summary는 유형 지정 필요, query_transactions는 집계 도구가 아님.\n\n"
        "[4. 필수 파라미터 추출]\n• 금융회사 번호 134 → bank_id=134",
    ),
    (
        19, "이상거래 통계 및 현황 조회",
        "금융회사 종합 현황 보고 (① 정답 도구)",
        "내부 위험 모니터링 시스템에서 특정 금융회사가 고위험으로 분류되어 상세 현황을 즉시 확인해야 하는 상황",
        "금융회사 77번 종합 보고서 내줘. 이상거래 유형 분포와 거래 규모도 포함해서.",
        "get_institution_report", [("bank_id", "77")],
        "특정 금융회사의 이상거래 유형 분포와 거래 규모를 포함한 종합 현황 조회가 목적. "
        "get_institution_report는 fraud_type_distribution과 total_count·total_amount를 포함하므로 적합. "
        "get_statistics는 전체 금융회사 통합 집계로 특정 기관 분리 불가. "
        "get_fraud_type_summary는 유형 코드를 하나씩 지정해야 하므로 전체 유형 분포 조회에 비효율적.",
        "[1. 질문 주제 파악]\n• 사용자는 고위험 금융회사 모니터링 담당 AML 직원이며, 금융회사 77번 종합 현황 조회가 주제임.\n\n"
        "[2. 질문 목적 분석]\n• 금융회사 77번의 이상거래 유형 분포와 거래 규모를 포함한 종합 보고서를 즉시 확인하고자 함.\n\n"
        "[3. 필요 기능 도출]\n• get_institution_report는 fraud_type_distribution과 거래 규모를 단일 호출로 반환. "
        "get_statistics는 전체 집계, get_fraud_type_summary는 유형별 개별 조회로 부적합.\n\n"
        "[4. 필수 파라미터 추출]\n• 금융회사 번호 77 → bank_id=77",
    ),

    # ── rank_risky_transactions ─────────────────────────────────────────
    (
        20, "AML 탐지·분석 및 보고서 작성",
        "위험도 상위 거래 일괄 랭킹 (① 정답 도구)",
        "야간 배치 처리 후 AI 모델 기준으로 당일 거래 중 가장 위험한 거래를 파악해야 하는 상황",
        "오늘 처리된 거래 중 이상거래 위험도가 가장 높은 20건을 찾아줘.",
        "rank_risky_transactions", [("top_k", "20")],
        "대량 거래에서 XGBoost 모델 예측 확률 기준 상위 K건을 자동으로 발굴하는 것이 목적. "
        "rank_risky_transactions는 일괄 예측 후 위험도 상위 top_k건을 반환하므로 적합. "
        "predict_fraud는 단건 예측 도구로 대량 거래 전체 랭킹 작업에 부적합. "
        "query_transactions는 확률 예측 기능이 없으므로 위험도 기준 정렬 불가.",
        "[1. 질문 주제 파악]\n• 사용자는 일일 배치 모니터링 담당 AML 직원이며, 위험도 상위 거래 자동 발굴이 주제임.\n\n"
        "[2. 질문 목적 분석]\n• 당일 처리 거래 중 AI 모델 기준 이상거래 확률이 가장 높은 20건을 파악하고자 함.\n\n"
        "[3. 필요 기능 도출]\n• rank_risky_transactions는 배치 예측 후 확률 내림차순으로 상위 K건을 반환하므로 적합. "
        "predict_fraud는 6개 속성을 직접 입력하는 단건 예측, query_transactions는 확률 예측 불가.\n\n"
        "[4. 필수 파라미터 추출]\n• 상위 20건 → top_k=20",
    ),
    (
        21, "AML 탐지·분석 및 보고서 작성",
        "위험도 상위 거래 일괄 랭킹 (① 정답 도구)",
        "주간 STR 검토 회의를 위해 최근 500건 거래 중 위험도 상위 10건을 선별하여 보고 자료를 준비해야 하는 상황",
        "거래 샘플 500건에서 이상거래 확률 상위 10건만 뽑아줘.",
        "rank_risky_transactions", [("sample_size", "500"), ("top_k", "10")],
        "지정한 샘플 크기(500건) 내에서 위험도 상위 K건(10건)을 발굴하는 것이 목적. "
        "rank_risky_transactions는 sample_size와 top_k를 모두 지원하여 적합. "
        "predict_fraud는 단건 입력이 필요하므로 500건 일괄 처리 불가. "
        "detect_aml_patterns는 그래프 패턴 탐지로 ML 확률 기반 랭킹과 무관.",
        "[1. 질문 주제 파악]\n• 사용자는 주간 회의 보고 자료를 준비하는 AML 직원이며, 샘플 내 위험도 상위 거래 선별이 주제임.\n\n"
        "[2. 질문 목적 분석]\n• 500건 샘플에서 XGBoost 예측 확률 기준 상위 10건을 추출하여 검토 자료를 만들고자 함.\n\n"
        "[3. 필요 기능 도출]\n• rank_risky_transactions(sample_size=500, top_k=10)이 샘플 범위와 반환 건수를 모두 제어. "
        "predict_fraud는 단건, detect_aml_patterns는 그래프 패턴으로 부적합.\n\n"
        "[4. 필수 파라미터 추출]\n• 샘플 크기 500 → sample_size=500\n• 상위 10건 → top_k=10",
    ),

    # ── generate_str (2nd) ───────────────────────────────────────────────
    (
        22, "AML 탐지·분석 및 보고서 작성",
        "의심거래보고서(STR) 공식 양식 작성 (① 정답 도구)",
        "다중 분석 도구를 활용한 종합 분석 완료 후 레이어링 의심 계좌에 대한 STR을 작성해야 하는 상황",
        "레이어링 패턴이 탐지된 계좌에 대해 STR을 작성해줘. 이상거래 확률 0.87, 탐지 패턴은 layering과 ring이야.",
        "generate_str",
        [("summary", "레이어링 및 순환거래 패턴이 탐지된 자금세탁 의심 계좌"),
         ("fraud_type", "자금세탁"), ("fraud_probability", "0.87"), ("aml_patterns", "['layering', 'ring']")],
        "레이어링·순환거래 패턴 탐지 결과와 ML 예측 확률을 종합하여 STR을 작성하는 것이 목적. "
        "generate_str는 fraud_probability와 aml_patterns 파라미터를 받아 "
        "의심 강도 산출 및 패턴 정보를 STR에 반영하므로 적합. "
        "detect_aml_patterns와 predict_fraud는 분석 도구로 보고서 작성 기능이 없음.",
        "[1. 질문 주제 파악]\n• 사용자는 종합 분석 결과를 취합하여 STR을 작성하는 AML 직원이며, "
        "다중 분석 결과를 통합한 STR 작성이 주제임.\n\n"
        "[2. 질문 목적 분석]\n• 이미 탐지된 layering/ring 패턴과 ML 확률 0.87을 근거로 자금세탁 STR을 작성하고자 함.\n\n"
        "[3. 필요 기능 도출]\n• generate_str는 fraud_probability(의심 강도 산출)와 aml_patterns(패턴 반영)를 "
        "모두 지원하므로 적합. detect_aml_patterns와 predict_fraud는 분석 도구이며 보고서 작성 기능 없음.\n\n"
        "[4. 필수 파라미터 추출]\n• summary=\"레이어링 및 순환거래 패턴이 탐지된 자금세탁 의심 계좌\"\n"
        "• fraud_type=\"자금세탁\"\n• 확률 0.87 → fraud_probability=0.87\n• 탐지 패턴 → aml_patterns=['layering', 'ring']",
    ),
]

# ---------------------------------------------------------------------------
# ③ 벤치마크 케이스 로드 (JSON → rows)
# ---------------------------------------------------------------------------

_CATEGORY_ORDER = [
    "cases_get_statistics", "cases_get_fraud_type_summary", "cases_get_account_profile",
    "cases_compare_periods", "cases_get_institution_report", "cases_query_transactions",
    "cases_analyze_network", "cases_detect_aml_patterns", "cases_rank_risky_transactions",
    "cases_predict_fraud", "cases_multi_tool",
    "cases_detect_ctr_candidates", "cases_score_account_risk", "cases_detect_monitoring_alerts",
    "cases_detect_dormant_reactivation", "cases_detect_smurfing_network",
    "cases_get_trend_analysis", "cases_analyze_channel_risk",
    "cases_get_receiving_account_profile", "cases_analyze_cross_institution_flow",
    "cases_missing_parameters",
    "cases_lookup_fiu_reference_types", "cases_validate_str_fields", "cases_get_aml_glossary",
]
_CATEGORY_LABELS = {k: k.replace("cases_", "") for k in _CATEGORY_ORDER}

_CASE_CAT_FILLS = [
    "EBF3FB", "E8F5E9", "FFF8E1", "FCE4D6", "F3E5F5",
    "E0F7FA", "FFF9C4", "E8EAF6", "FFEBEE", "E0F2F1", "F9FBE7",
    "DAEEF3", "FDE9D9", "E4DFEC", "D5E8D4", "DCE6F1",
    "FFF2CC", "F2DCDB", "D6E3F8", "E2EFDA", "F5F5F5",
    "E8E8E8", "DEDEDE", "CECECE",
]
_IRR_FONT = Font(color="9E9E9E", italic=True)


def _load_cases(dataset_dir: Path) -> list[dict]:
    rows = []
    for key in _CATEGORY_ORDER:
        path = dataset_dir / f"{key}.json"
        if not path.exists():
            continue
        cases = json.loads(path.read_text(encoding="utf-8"))
        label = _CATEGORY_LABELS[key]
        for case in cases:
            exp = case.get("expected", {})
            primary = exp.get("primary_tool", "")
            must = exp.get("tools_must_include", [])
            order = exp.get("tool_order", [])
            checks = exp.get("param_checks", {})
            rows.append({
                "category":    label,
                "id":          case.get("id", ""),
                "difficulty":  case.get("difficulty", ""),
                "type":        "irrelevance" if primary == "" else "normal",
                "question":    case.get("question", ""),
                "primary_tool": primary if primary else "-",
                "tools_must":  ", ".join(must) if must else "-",
                "tool_order":  " → ".join(order) if order else "-",
                "param_checks": json.dumps(checks, ensure_ascii=False) if checks else "-",
                "note":        case.get("note", ""),
            })
    return rows


# ---------------------------------------------------------------------------
# 시트 작성 함수
# ---------------------------------------------------------------------------

def _write_tool_sheet(wb: openpyxl.Workbook) -> None:
    ws = wb.active
    ws.title = "도구 정의"

    headers_row1 = ["", "sub domain", "task", "function", "", "parameters", "", "", ""]
    headers_row2 = ["", "", "", "name", "description", "required", "properties", "type", "description"]
    for col, val in enumerate(headers_row1, 1):
        c = ws.cell(row=1, column=col, value=val)
        c.font = _HEADER_FONT; c.fill = _HEADER_FILL; c.border = _BORDER; c.alignment = _WRAP_TOP
    for col, val in enumerate(headers_row2, 1):
        c = ws.cell(row=2, column=col, value=val)
        c.font = _HEADER_FONT; c.fill = _SUB_HEADER_FILL; c.border = _BORDER; c.alignment = _WRAP_TOP

    current_row = 3
    for tool_idx, (sub_domain, task, func_name, func_desc, params) in enumerate(_TOOLS):
        for p_idx, (req, prop_name, prop_type, prop_desc) in enumerate(params):
            row = current_row + p_idx
            if p_idx == 0:
                ws.cell(row=row, column=1, value=tool_idx + 1)
                ws.cell(row=row, column=2, value=sub_domain)
                ws.cell(row=row, column=3, value=task)
                ws.cell(row=row, column=4, value=func_name)
                ws.cell(row=row, column=5, value=func_desc)
            ws.cell(row=row, column=6, value=req)
            ws.cell(row=row, column=7, value=prop_name)
            ws.cell(row=row, column=8, value=prop_type)
            ws.cell(row=row, column=9, value=prop_desc)
        current_row += len(params)
        ws.cell(row=current_row, column=6, value="-")
        ws.cell(row=current_row, column=7, value="-")
        ws.cell(row=current_row, column=8, value="-")
        current_row += 1

    for row in ws.iter_rows(min_row=3, max_row=current_row - 1, max_col=9):
        for cell in row:
            cell.border = _BORDER; cell.alignment = _WRAP_TOP

    for col, width in zip("ABCDEFGHI", [5, 26, 40, 30, 60, 10, 26, 10, 60]):
        ws.column_dimensions[col].width = width


def _write_scenario_sheet(wb: openpyxl.Workbook) -> None:
    ws = wb.create_sheet("평가 시나리오")

    headers_row1 = ["", "sub domain", "task", "시나리오 설정", "사용자 질문",
                    "정답", "", "", "정답 설정 근거", "모델 사고 과정"]
    headers_row2 = ["", "", "", "", "", "function.name", "properties", "properties 입력값", "", ""]
    for col, val in enumerate(headers_row1, 1):
        c = ws.cell(row=1, column=col, value=val)
        c.font = _HEADER_FONT; c.fill = _HEADER_FILL; c.border = _BORDER; c.alignment = _WRAP_TOP
    for col, val in enumerate(headers_row2, 1):
        c = ws.cell(row=2, column=col, value=val)
        c.font = Font(bold=True, color="FFFFFF", size=11)
        c.fill = _SUB_HEADER_FILL; c.border = _BORDER; c.alignment = _WRAP_TOP

    for current_row, sc in enumerate(_SCENARIOS, start=3):
        no, sub_domain, task, scenario_setting, user_q, func_name, params, rationale, reasoning = sc
        props_text = "\n".join(p for p, _ in params) if params else "(없음)"
        vals_text = "\n".join(v for _, v in params) if params else "(없음)"
        for col, val in enumerate(
            [no, sub_domain, task, scenario_setting, user_q,
             func_name or "(없음)", props_text, vals_text, rationale, reasoning], 1
        ):
            c = ws.cell(row=current_row, column=col, value=val)
            c.border = _BORDER
            c.alignment = _WRAP_TOP if col in (4, 5, 7, 8, 9, 10) else _CENTER_MID

    for row in ws.iter_rows(min_row=3, max_row=len(_SCENARIOS) + 2, max_col=10):
        for cell in row:
            cell.border = _BORDER

    for col, width in zip("ABCDEFGHIJ", [5, 26, 40, 45, 50, 26, 22, 35, 55, 65]):
        ws.column_dimensions[col].width = width


def _write_case_sheets(wb: openpyxl.Workbook, rows: list[dict]) -> None:
    # 케이스 전체 시트
    ws = wb.create_sheet("벤치마크 케이스")
    headers = ["category", "id", "difficulty", "type", "question",
               "primary_tool", "tools_must_include", "tool_order", "param_checks", "note"]
    for col, h in enumerate(headers, 1):
        c = ws.cell(row=1, column=col, value=h)
        c.font = _HEADER_FONT; c.fill = _HEADER_FILL; c.border = _BORDER; c.alignment = _CENTER_MID

    cat_color: dict[str, str] = {}
    color_idx = 0
    for row_idx, r in enumerate(rows, 2):
        cat = r["category"]
        if cat not in cat_color:
            cat_color[cat] = _CASE_CAT_FILLS[color_idx % len(_CASE_CAT_FILLS)]
            color_idx += 1
        fill = PatternFill(start_color=cat_color[cat], end_color=cat_color[cat], fill_type="solid")
        is_irr = r["type"] == "irrelevance"
        for col, val in enumerate(
            [r["category"], r["id"], r["difficulty"], r["type"], r["question"],
             r["primary_tool"], r["tools_must"], r["tool_order"], r["param_checks"], r["note"]], 1
        ):
            c = ws.cell(row=row_idx, column=col, value=val)
            c.border = _BORDER; c.fill = fill
            c.font = _IRR_FONT if is_irr else Font()
            c.alignment = _WRAP_TOP if col in (5, 9, 10) else _CENTER_MID

    for col, width in zip("ABCDEFGHIJ", [22, 18, 11, 13, 55, 26, 40, 38, 50, 40]):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A2"

    # 케이스 요약 시트
    ws_sum = wb.create_sheet("케이스 요약")
    sum_headers = ["category", "total", "normal", "irrelevance", "easy", "medium", "hard"]
    for col, h in enumerate(sum_headers, 1):
        c = ws_sum.cell(row=1, column=col, value=h)
        c.font = _HEADER_FONT; c.fill = _HEADER_FILL; c.border = _BORDER; c.alignment = _CENTER_MID

    stats: dict[str, dict] = {}
    for r in rows:
        cat = r["category"]
        if cat not in stats:
            stats[cat] = {"total": 0, "normal": 0, "irrelevance": 0, "easy": 0, "medium": 0, "hard": 0}
        stats[cat]["total"] += 1
        stats[cat][r["type"]] += 1
        d = r["difficulty"]
        if d in stats[cat]:
            stats[cat][d] += 1

    total_sum = {"total": 0, "normal": 0, "irrelevance": 0, "easy": 0, "medium": 0, "hard": 0}
    for row_idx, key in enumerate(_CATEGORY_ORDER, 2):
        label = _CATEGORY_LABELS[key]
        if label not in stats:
            continue
        s = stats[label]
        fill_color = _CASE_CAT_FILLS[(row_idx - 2) % len(_CASE_CAT_FILLS)]
        fill = PatternFill(start_color=fill_color, end_color=fill_color, fill_type="solid")
        for col, val in enumerate(
            [label, s["total"], s["normal"], s["irrelevance"], s["easy"], s["medium"], s["hard"]], 1
        ):
            c = ws_sum.cell(row=row_idx, column=col, value=val)
            c.border = _BORDER; c.fill = fill; c.alignment = _CENTER_MID
        for k in total_sum:
            total_sum[k] += s[k]

    total_row_idx = len(_CATEGORY_ORDER) + 2
    bold_fill = PatternFill(start_color="D6E4F0", end_color="D6E4F0", fill_type="solid")
    for col, val in enumerate(
        ["TOTAL", total_sum["total"], total_sum["normal"], total_sum["irrelevance"],
         total_sum["easy"], total_sum["medium"], total_sum["hard"]], 1
    ):
        c = ws_sum.cell(row=total_row_idx, column=col, value=val)
        c.font = _BOLD; c.border = _BORDER; c.fill = bold_fill; c.alignment = _CENTER_MID

    for col in range(1, 8):
        ws_sum.column_dimensions[ws_sum.cell(row=1, column=col).column_letter].width = 22


# ---------------------------------------------------------------------------
# main
# ---------------------------------------------------------------------------

if __name__ == "__main__":
    dataset_dir = Path(__file__).parent / "dataset"
    output_path = Path(__file__).parent / "benchmark_dataset.xlsx"

    cases = _load_cases(dataset_dir)

    wb = openpyxl.Workbook()
    _write_tool_sheet(wb)
    _write_scenario_sheet(wb)
    _write_case_sheets(wb, cases)
    wb.save(output_path)

    normal = sum(1 for r in cases if r["type"] == "normal")
    irr = sum(1 for r in cases if r["type"] == "irrelevance")
    print(f"Saved: {output_path}")
    print(f"  [1] tool definitions : {len(_TOOLS)}")
    print(f"  [2] eval scenarios   : {len(_SCENARIOS)}")
    print(f"  [3] bench cases      : {len(cases)} (normal={normal}, irrelevance={irr})")
    print(f"  [4] case summary     : by category")
